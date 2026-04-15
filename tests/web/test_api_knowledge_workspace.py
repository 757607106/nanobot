from __future__ import annotations

import io
import time
from pathlib import Path

from fastapi.testclient import TestClient

from nanobot.config import loader as config_loader
from nanobot.config.loader import save_config
from nanobot.config.schema import Config
from nanobot.platform.knowledge.preview_artifacts import KnowledgePreviewArtifacts
from nanobot.web.app import create_app
from tests.knowledge_test_utils import FakeRAGEngine


def _make_test_config(tmp_path, monkeypatch) -> Config:
    # Keep instance_id unique across pytest runs and test cases.
    config_path = tmp_path / f"{tmp_path.parent.name}-{tmp_path.name}" / "config.json"
    config_path.parent.mkdir(parents=True, exist_ok=True)
    workspace = tmp_path / "workspace"
    config = Config()
    config.agents.defaults.workspace = str(workspace)
    save_config(config, config_path)
    monkeypatch.setattr(config_loader, "_current_config_path", config_path)
    return config


def _bootstrap_admin(client: TestClient) -> None:
    response = client.post(
        "/api/v1/auth/bootstrap",
        json={"username": "admin", "password": "bootstrap-pass-123"},
    )
    assert response.status_code == 201

def _wait_for_job(client: TestClient, kb_id: str, job_id: str) -> dict:
    deadline = time.time() + 5.0
    while time.time() < deadline:
        response = client.get(f"/api/v1/knowledge-bases/{kb_id}/jobs")
        assert response.status_code == 200
        job = next((item for item in response.json()["data"] if item["jobId"] == job_id), None)
        if job and job["status"] in {"succeeded", "failed"}:
            return job
        time.sleep(0.05)
    raise AssertionError(f"Timed out waiting for job {job_id}")


def _wait_for_evaluation(client: TestClient, kb_id: str, task_id: str) -> dict:
    deadline = time.time() + 5.0
    while time.time() < deadline:
        response = client.get(f"/api/v1/knowledge-bases/{kb_id}/evaluation/results/{task_id}")
        assert response.status_code == 200
        result = response.json()["data"]
        if result["status"] in {"completed", "failed"}:
            return result
        time.sleep(0.05)
    raise AssertionError(f"Timed out waiting for evaluation {task_id}")


def _build_docx_bytes(*paragraphs: str) -> bytes:
    from docx import Document

    buffer = io.BytesIO()
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(buffer)
    return buffer.getvalue()


def _build_xlsx_bytes(*rows: tuple[object, ...]) -> bytes:
    from openpyxl import Workbook

    buffer = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(list(row))
    workbook.save(buffer)
    return buffer.getvalue()


def test_web_api_unified_workspace_flow(tmp_path, monkeypatch) -> None:
    config = _make_test_config(tmp_path, monkeypatch)
    monkeypatch.setattr("nanobot.web.app.create_rag_engine_from_config", lambda config, instance_dir: FakeRAGEngine())
    monkeypatch.setattr(
        "nanobot.web.runtime_services.config.create_rag_engine_from_config",
        lambda config, instance_dir: FakeRAGEngine(),
    )
    convert_calls = {"count": 0}

    def _fake_convert(_source_path: Path, _preview_dir: Path, target_path: Path):
        convert_calls["count"] += 1
        target_path.write_bytes(b"%PDF-1.7\n%nanobot-preview\n")
        return target_path, None

    monkeypatch.setattr(
        KnowledgePreviewArtifacts,
        "_convert_presentation_to_pdf",
        staticmethod(_fake_convert),
    )

    with TestClient(create_app(config, static_dir=tmp_path / "missing-static")) as client:
        _bootstrap_admin(client)
        client.app.state.knowledge._generate_with_llm = lambda **kwargs: None  # type: ignore[method-assign]

        created = client.post(
            "/api/v1/knowledge-bases",
            json={
                "name": "Ops API KB",
                "kbType": "milvus",
                "description": "API smoke test",
            },
        )
        assert created.status_code == 201
        kb_id = created.json()["data"]["kbId"]
        assert created.json()["data"]["kbType"] == "lightrag"
        assert created.json()["data"]["query_params"]["mode"] == "mix"

        unsupported = client.post(
            "/api/v1/knowledge-bases",
            json={
                "name": "Unsupported KB",
                "kbType": "dify",
                "description": "Should fail",
            },
        )
        assert unsupported.status_code == 400

        schema = client.get(f"/api/v1/knowledge-bases/{kb_id}/query-params/schema")
        assert schema.status_code == 200
        assert schema.json()["data"]["type"] == "lightrag"

        description = client.post(
            "/api/v1/knowledge-bases/generate-description",
            json={"name": "Ops API KB", "currentDescription": "", "fileList": ["/runbook.md"]},
        )
        assert description.status_code == 200
        assert description.json()["data"]["description"]

        uploaded = client.post(
            f"/api/v1/knowledge-bases/{kb_id}/files",
            files={
                "file": (
                    "runbook.md",
                    b"# API Runbook\n\nRestart the worker after draining the queue.\n\nClear the token cache.\n",
                    "text/markdown",
                )
            },
        )
        assert uploaded.status_code == 201
        file_id = uploaded.json()["data"]["items"][0]["fileId"]

        guide_upload = client.post(
            f"/api/v1/knowledge-bases/{kb_id}/files",
            files=[
                ("file", ("guide.md", b"![diagram](diagram.png)\n\n[External](https://example.com/docs)\n", "text/markdown")),
                ("file", ("diagram.png", b"\x89PNG\r\n\x1a\ndiagram-bytes", "image/png")),
                ("file", ("page.html", b"<html><body><img src='diagram.png'><a href='guide.md'>Guide</a></body></html>", "text/html")),
            ],
        )
        assert guide_upload.status_code == 201
        guide_file_id = next(item["fileId"] for item in guide_upload.json()["data"]["items"] if item["filename"] == "guide.md")
        diagram_file_id = next(item["fileId"] for item in guide_upload.json()["data"]["items"] if item["filename"] == "diagram.png")
        html_file_id = next(item["fileId"] for item in guide_upload.json()["data"]["items"] if item["filename"] == "page.html")

        guide_preview = client.get(f"/api/v1/knowledge-bases/{kb_id}/files/{guide_file_id}/preview")
        assert guide_preview.status_code == 200
        guide_preview_payload = guide_preview.json()["data"]
        assert guide_preview_payload["previewKind"] == "markdown"
        assert guide_preview_payload["baseUrl"].endswith(f"/api/v1/knowledge-bases/{kb_id}/files/{guide_file_id}/preview/assets/")

        guide_inline = client.get(
            f"/api/v1/knowledge-bases/{kb_id}/files/{guide_file_id}/download",
            params={"variant": "raw", "disposition": "inline"},
        )
        assert guide_inline.status_code == 200
        assert guide_inline.headers["content-disposition"].startswith("inline;")

        guide_asset = client.get(
            f"/api/v1/knowledge-bases/{kb_id}/files/{guide_file_id}/preview/assets/diagram.png"
        )
        assert guide_asset.status_code == 200
        assert guide_asset.headers["content-type"].startswith("image/png")
        assert guide_asset.headers["content-disposition"].startswith("inline;")

        html_preview = client.get(f"/api/v1/knowledge-bases/{kb_id}/files/{html_file_id}/preview")
        assert html_preview.status_code == 200
        assert html_preview.json()["data"]["previewKind"] == "html"

        html_inline = client.get(
            f"/api/v1/knowledge-bases/{kb_id}/files/{html_file_id}/download",
            params={"variant": "raw", "disposition": "inline"},
        )
        assert html_inline.status_code == 200
        assert html_inline.headers["content-disposition"].startswith("inline;")
        assert html_inline.headers["content-security-policy"].startswith("sandbox;")

        office_upload = client.post(
            f"/api/v1/knowledge-bases/{kb_id}/files",
            files=[
                (
                    "file",
                    (
                        "manual.docx",
                        _build_docx_bytes("D9 操作手册", "按步骤执行即可。"),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    ),
                ),
                (
                    "file",
                    (
                        "metrics.xlsx",
                        _build_xlsx_bytes(("Name", "Value"), ("CPU", 42)),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    ),
                ),
                (
                    "file",
                    (
                        "deck.pptx",
                        b"pptx-placeholder",
                        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    ),
                ),
            ],
        )
        assert office_upload.status_code == 201
        office_docx_id = next(item["fileId"] for item in office_upload.json()["data"]["items"] if item["filename"] == "manual.docx")
        office_xlsx_id = next(item["fileId"] for item in office_upload.json()["data"]["items"] if item["filename"] == "metrics.xlsx")
        office_pptx_id = next(item["fileId"] for item in office_upload.json()["data"]["items"] if item["filename"] == "deck.pptx")

        office_docx_preview = client.get(f"/api/v1/knowledge-bases/{kb_id}/files/{office_docx_id}/preview")
        assert office_docx_preview.status_code == 200
        assert office_docx_preview.json()["data"]["previewKind"] == "html"
        assert office_docx_preview.json()["data"]["contentType"] == "text/html"

        office_xlsx_preview = client.get(f"/api/v1/knowledge-bases/{kb_id}/files/{office_xlsx_id}/preview")
        assert office_xlsx_preview.status_code == 200
        assert office_xlsx_preview.json()["data"]["previewKind"] == "html"
        assert office_xlsx_preview.json()["data"]["contentType"] == "text/html"

        office_pptx_preview = client.get(f"/api/v1/knowledge-bases/{kb_id}/files/{office_pptx_id}/preview")
        assert office_pptx_preview.status_code == 200
        assert office_pptx_preview.json()["data"]["previewKind"] == "pdf"
        assert office_pptx_preview.json()["data"]["contentType"] == "application/pdf"

        office_docx_inline = client.get(
            f"/api/v1/knowledge-bases/{kb_id}/files/{office_docx_id}/download",
            params={"variant": "preview", "disposition": "inline"},
        )
        assert office_docx_inline.status_code == 200
        assert office_docx_inline.headers["content-type"].startswith("text/html")

        office_xlsx_inline = client.get(
            f"/api/v1/knowledge-bases/{kb_id}/files/{office_xlsx_id}/download",
            params={"variant": "preview", "disposition": "inline"},
        )
        assert office_xlsx_inline.status_code == 200
        assert office_xlsx_inline.headers["content-type"].startswith("text/html")

        office_pptx_inline = client.get(
            f"/api/v1/knowledge-bases/{kb_id}/files/{office_pptx_id}/download",
            params={"variant": "preview", "disposition": "inline"},
        )
        assert office_pptx_inline.status_code == 200
        assert office_pptx_inline.headers["content-type"].startswith("application/pdf")
        assert convert_calls["count"] == 1

        parse_missing_ids = client.post(f"/api/v1/knowledge-bases/{kb_id}/files/parse", json={})
        assert parse_missing_ids.status_code == 400

        index_missing_ids = client.post(f"/api/v1/knowledge-bases/{kb_id}/files/index", json={})
        assert index_missing_ids.status_code == 400

        parsed = client.post(f"/api/v1/knowledge-bases/{kb_id}/files/parse", json={"file_ids": [file_id]})
        assert parsed.status_code == 202
        assert _wait_for_job(client, kb_id, parsed.json()["data"]["job"]["jobId"])["status"] == "succeeded"

        indexed = client.post(f"/api/v1/knowledge-bases/{kb_id}/files/index", json={"file_ids": [file_id]})
        assert indexed.status_code == 202
        assert _wait_for_job(client, kb_id, indexed.json()["data"]["job"]["jobId"])["status"] == "succeeded"

        queried = client.post(
            f"/api/v1/knowledge-bases/{kb_id}/query",
            json={
                "query": "How do we restart the worker?",
                "mode": "mix",
                "top_k": 4,
                "only_need_context": False,
            },
        )
        assert queried.status_code == 200
        query_payload = queried.json()["data"]
        assert query_payload["metadata"]["kbType"] == "lightrag"
        assert query_payload["data"]["chunks"]

        multimodal_query = client.post(
            f"/api/v1/knowledge-bases/{kb_id}/query",
            json={
                "query": "Describe the uploaded diagram",
                "mode": "mix",
                "only_need_context": False,
                "multimodal_content": [
                    {
                        "type": "image",
                        "file_id": diagram_file_id,
                    }
                ],
            },
        )
        assert multimodal_query.status_code == 200
        multimodal_payload = multimodal_query.json()["data"]
        assert multimodal_payload["metadata"]["multimodalQuery"] is True
        assert multimodal_payload["metadata"]["multimodalContentTypes"] == ["image"]
        assert multimodal_payload["data"]["chunks"] == []
        assert "multimodal[image]" in str(multimodal_payload.get("message") or "")

        faq_source = client.post(
            f"/api/v1/knowledge-bases/{kb_id}/sources",
            json={
                "sourceType": "faq_table",
                "title": "FAQ Source",
                "items": [
                    {
                        "question": "How do we clear the token cache?",
                        "answer": "Drain the queue, then clear the token cache before resuming traffic.",
                    }
                ],
            },
        )
        assert faq_source.status_code == 201
        faq_file_id = faq_source.json()["data"]["fileId"]

        faq_parsed = client.post(f"/api/v1/knowledge-bases/{kb_id}/files/parse", json={"file_ids": [faq_file_id]})
        assert faq_parsed.status_code == 202
        assert _wait_for_job(client, kb_id, faq_parsed.json()["data"]["job"]["jobId"])["status"] == "succeeded"

        faq_indexed = client.post(f"/api/v1/knowledge-bases/{kb_id}/files/index", json={"file_ids": [faq_file_id]})
        assert faq_indexed.status_code == 202
        assert _wait_for_job(client, kb_id, faq_indexed.json()["data"]["job"]["jobId"])["status"] == "succeeded"

        qa_upload = client.post(
            f"/api/v1/knowledge-bases/{kb_id}/files",
            files={
                "file": (
                    "faq.md",
                    (
                        b"Q: How do we drain the queue?\n"
                        b"A: Pause intake and wait for workers to finish.\n"
                        b"Q: How do we clear the cache?\n"
                        b"A: Run the cache reset task before resuming traffic.\n"
                    ),
                    "text/markdown",
                )
            },
        )
        assert qa_upload.status_code == 201
        qa_file_id = qa_upload.json()["data"]["items"][0]["fileId"]

        qa_parsed = client.post(f"/api/v1/knowledge-bases/{kb_id}/files/parse", json={"file_ids": [qa_file_id]})
        assert qa_parsed.status_code == 202
        assert _wait_for_job(client, kb_id, qa_parsed.json()["data"]["job"]["jobId"])["status"] == "succeeded"

        qa_indexed = client.post(
            f"/api/v1/knowledge-bases/{kb_id}/files/index",
            json={
                "file_ids": [qa_file_id],
                "params": {
                    "chunk_preset_id": "qa",
                    "chunk_size": 400,
                    "chunk_overlap": 50,
                },
            },
        )
        assert qa_indexed.status_code == 202
        assert _wait_for_job(client, kb_id, qa_indexed.json()["data"]["job"]["jobId"])["status"] == "succeeded"

        files_after_index = client.get(f"/api/v1/knowledge-bases/{kb_id}/files")
        assert files_after_index.status_code == 200
        qa_file = next(item for item in files_after_index.json()["data"]["items"] if item["fileId"] == qa_file_id)
        assert qa_file["processingParams"]["chunk_preset_id"] == "qa"

        qa_detail = client.get(f"/api/v1/knowledge-bases/{kb_id}/files/{qa_file_id}/detail")
        assert qa_detail.status_code == 200
        detail_payload = qa_detail.json()["data"]
        assert "How do we drain the queue?" in detail_payload["content"]
        assert len(detail_payload["chunks"]) == 2

        updated_query_params = client.put(
            f"/api/v1/knowledge-bases/{kb_id}/query-params",
            json={
                "chunk_top_k": 8,
                "enable_rerank": True,
            },
        )
        assert updated_query_params.status_code == 200
        query_params_payload = updated_query_params.json()["data"]
        assert query_params_payload["mode"] == "mix"
        assert query_params_payload["chunk_top_k"] == 8
        assert query_params_payload["enable_rerank"] is True

        filtered_query = client.post(
            f"/api/v1/knowledge-bases/{kb_id}/query",
            json={
                "query": "How do we clear the cache?",
                "mode": "mix",
                "top_k": 2,
                "file_ids": [qa_file_id],
            },
        )
        assert filtered_query.status_code == 200
        filtered_payload = filtered_query.json()["data"]
        assert filtered_payload["metadata"]["fileFilterApplied"] is True
        assert any("cache reset task" in str(item.get("content") or "").lower() for item in filtered_payload["data"]["chunks"])

        generated = client.post(
            f"/api/v1/knowledge-bases/{kb_id}/benchmarks/generate",
            json={"count": 1, "name": "API Benchmark"},
        )
        assert generated.status_code == 201
        benchmark_id = generated.json()["data"]["benchmarkId"]

        listed = client.get(f"/api/v1/knowledge-bases/{kb_id}/benchmarks")
        assert listed.status_code == 200
        assert listed.json()["data"][0]["benchmarkId"] == benchmark_id

        evaluation = client.post(
            f"/api/v1/knowledge-bases/{kb_id}/evaluation/run",
            json={"benchmarkId": benchmark_id},
        )
        assert evaluation.status_code == 202
        task_id = evaluation.json()["data"]["taskId"]

        result = _wait_for_evaluation(client, kb_id, task_id)
        assert result["status"] == "completed"
        assert result["details"]


def test_web_api_legacy_routes_removed(tmp_path, monkeypatch) -> None:
    config = _make_test_config(tmp_path, monkeypatch)
    monkeypatch.setattr("nanobot.web.app.create_rag_engine_from_config", lambda config, instance_dir: FakeRAGEngine())
    monkeypatch.setattr(
        "nanobot.web.runtime_services.config.create_rag_engine_from_config",
        lambda config, instance_dir: FakeRAGEngine(),
    )

    with TestClient(create_app(config, static_dir=tmp_path / "missing-static")) as client:
        _bootstrap_admin(client)
        assert client.post("/api/knowledge/databases", json={"name": "legacy"}).status_code == 404
        assert client.get("/api/graph/list").status_code == 404
        assert client.post("/api/knowledge/files/upload?db_id=any").status_code == 404
